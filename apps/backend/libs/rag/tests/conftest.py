import io
import json

import docx
import httpx
import pymupdf


def make_pdf(
    pages: list[str], *, title: str | None = None, toc: list[list[object]] | None = None
) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=10)
    if title:
        doc.set_metadata({"title": title})
    if toc:
        doc.set_toc(toc)
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    document = docx.Document()
    document.core_properties.title = "Handbook"
    document.add_heading("Onboarding", level=1)
    document.add_paragraph("Welcome to the team. Laptops are issued on day one.")
    document.add_heading("Security", level=2)
    document.add_paragraph("Enable two-factor authentication on every account.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Tool", "Owner"
    table.cell(1, 0).text, table.cell(1, 1).text = "VPN", "IT"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def fake_embedding(text: str, dim: int = 8) -> list[float]:
    return [float((hash(text) >> (i * 4)) % 7 + 1) for i in range(dim)]


class FakeGateway:
    """Handles /v1/embeddings and /v1/chat/completions (metadata extraction)."""

    def __init__(self) -> None:
        self.embed_calls: list[list[str]] = []
        self.chat_calls = 0
        self.embed_status = 200
        self.chat_reply = json.dumps({"summary": "A short summary.", "keywords": ["alpha", "beta"]})

    def handler(self, request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        if request.url.path == "/v1/embeddings":
            if self.embed_status != 200:
                return httpx.Response(
                    self.embed_status, json={"error": {"code": "x", "message": "embed failed"}}
                )
            inputs = body["input"] if isinstance(body["input"], list) else [body["input"]]
            self.embed_calls.append(inputs)
            return httpx.Response(
                200,
                json={
                    "data": [
                        {"index": i, "embedding": fake_embedding(t)} for i, t in enumerate(inputs)
                    ]
                },
            )
        self.chat_calls += 1
        return httpx.Response(
            200, json={"choices": [{"message": {"role": "assistant", "content": self.chat_reply}}]}
        )
