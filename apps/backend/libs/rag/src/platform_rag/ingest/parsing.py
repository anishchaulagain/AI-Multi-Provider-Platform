"""Bytes -> text sections with page and heading context.

PDFs keep page boundaries (and the outline as section paths); DOCX and HTML are
converted to markdown so headings drive chunking just like native markdown.
"""

import io
import re
from dataclasses import dataclass, field
from typing import Any

from platform_rag.formats import DocumentParseError, Format, decode_text, detect_format


@dataclass
class Section:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class ParsedDocument:
    format: Format
    title: str | None
    sections: list[Section]
    # True when the text is markdown (headings should drive chunking).
    is_markdown: bool = False
    page_count: int | None = None
    metadata: dict[str, str] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections)


def parse(data: bytes, filename: str) -> ParsedDocument:
    fmt = detect_format(filename, data)
    parser = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "html": _parse_html,
        "markdown": _parse_markdown,
        "text": _parse_text,
    }[fmt]
    try:
        doc = parser(data)
    except DocumentParseError:
        raise
    except Exception as exc:  # parser libraries raise a zoo of exception types
        raise DocumentParseError(f"Could not read {fmt} file: {exc}") from exc

    doc.sections = [s for s in doc.sections if s.text.strip()]
    if not doc.sections:
        hint = " (scanned PDFs need OCR, which is not supported yet)" if fmt == "pdf" else ""
        raise DocumentParseError(f"No extractable text{hint}")
    doc.title = (doc.title or "").strip() or _title_from_text(doc.text) or filename
    return doc


def _title_from_text(text: str) -> str | None:
    for line in text.splitlines():
        line = line.strip().lstrip("#").strip()
        if 3 <= len(line) <= 200:
            return line
    return None


# ---------- PDF ----------
def _parse_pdf(data: bytes) -> ParsedDocument:
    import pymupdf

    try:
        pdf: Any = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentParseError(f"Corrupt or unreadable PDF: {exc}") from exc
    with pdf:
        if pdf.needs_pass:
            raise DocumentParseError("PDF is password-protected")
        # Outline entries: [level, title, page]. Track the heading path in effect per page.
        toc = sorted(pdf.get_toc(simple=True), key=lambda e: e[2])
        sections: list[Section] = []
        path: list[str] = []
        toc_i = 0
        for index, page in enumerate(pdf):
            number = index + 1
            while toc_i < len(toc) and toc[toc_i][2] <= number:
                level, heading = int(toc[toc_i][0]), str(toc[toc_i][1]).strip()
                path = [*path[: level - 1], heading]
                toc_i += 1
            text = page.get_text("text", sort=True)
            sections.append(
                Section(text=_clean(text), page=number, section=" > ".join(path) or None)
            )
        meta = {k: str(v) for k, v in (pdf.metadata or {}).items() if v}
        return ParsedDocument(
            format="pdf",
            title=meta.get("title"),
            sections=sections,
            page_count=pdf.page_count,
            metadata={k: meta[k] for k in ("author", "subject") if k in meta},
        )


# ---------- DOCX ----------
def _parse_docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    title = document.core_properties.title or None
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style == "Title":
            title = title or text
            lines.append(f"# {text}")
        elif match := re.match(r"Heading (\d)", style):
            lines.append(f"{'#' * min(int(match.group(1)), 6)} {text}")
        elif style.startswith("List"):
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")
    return ParsedDocument(
        format="docx", title=title, sections=[Section("\n".join(lines))], is_markdown=True
    )


# ---------- HTML ----------
_BLOCKS = [
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "p",
    "li",
    "pre",
    "blockquote",
    "td",
    "th",
    "dt",
    "dd",
]


def _parse_html(data: bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(decode_text(data), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["script", "style", "noscript", "template", "svg", "nav", "footer"]):
        tag.decompose()
    root = soup.body or soup
    lines: list[str] = []
    for el in root.find_all(_BLOCKS):
        if el.find(_BLOCKS):  # only leaf blocks, so nested text isn't repeated
            continue
        text = " ".join(el.get_text(" ", strip=True).split())
        if not text:
            continue
        if el.name and el.name[0] == "h" and el.name[1:].isdigit():
            lines.append(f"{'#' * int(el.name[1:])} {text}")
        elif el.name == "li":
            lines.append(f"- {text}")
        elif el.name == "pre":
            lines.append(f"```\n{el.get_text()}\n```")
        else:
            lines.append(text)
        lines.append("")
    if not lines:
        lines = [root.get_text("\n", strip=True)]
    return ParsedDocument(
        format="html", title=title, sections=[Section("\n".join(lines))], is_markdown=True
    )


# ---------- Markdown / text ----------
def _parse_markdown(data: bytes) -> ParsedDocument:
    text = decode_text(data)
    return ParsedDocument(format="markdown", title=None, sections=[Section(text)], is_markdown=True)


def _parse_text(data: bytes) -> ParsedDocument:
    return ParsedDocument(format="text", title=None, sections=[Section(_clean(decode_text(data)))])


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()
